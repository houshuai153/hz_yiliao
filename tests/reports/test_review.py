"""报告人工审核流程：状态机、修订与溯源、通过前检查、导出门禁、持久化与旧数据兼容。"""
import io

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import select, text

from app.audit.models import AuditLog
from app.db.models import ReportDraft
from app.main import create_app
from app.reports.review import REQUIRED_CONFIRMATIONS
from app.settings import Settings

ALL_CONFIRMS = list(REQUIRED_CONFIRMATIONS)


@pytest.fixture
def h(auth):
    return {**auth, "X-Operator": "reviewer-a"}


def new_draft(api, h, pid="PRJ-SM-001", kind="validation", runs=()):
    r = api.post("/api/reports/draft", headers=h, json={"project_id": pid, "run_ids": list(runs), "kind": kind})
    assert r.status_code == 200, r.text
    return {**r.json(), **detail(api, h, r.json()["draft_id"])}       # 生成响应 + 预览内容


def detail(api, h, did, **params):
    r = api.get(f"/api/reports/{did}", headers=h, params=params)
    assert r.status_code == 200, r.text
    return r.json()


def blocks(d):
    return [dict(b, section=s["title"]) for s in d["sections"] for b in s["blocks"]]


def pick(d, *, kind, tagged=None, with_digits=None):
    for b in blocks(d):
        if b["type"] != "sentence" or b["kind"] != kind or b["edited"]:
            continue
        if any(t in b["section"] for t in ("附录", "签名")):
            continue
        if tagged is not None and bool(b["tags"]) != tagged:
            continue
        if with_digits is not None and any(c.isdigit() for c in b["text"]) != with_digits:
            continue
        return b
    raise AssertionError(f"没有找到 {kind} 句")


def act(api, h, did, action, **kw):
    return api.post(f"/api/reports/{did}/transition", headers=h, json={"action": action, **kw})


def approve_kw(d):
    return {"comment": "已逐项核对业务结论与来源", "confirmations": ALL_CONFIRMS,
            "missing_ack": {f: "保留为待补充，待业务提供原始记录后重新生成" for f in d["missing_fields"]}}


def to_in_review(api, h, did):
    assert act(api, h, did, "start_review").status_code == 200


# ---------- 1. 新建草稿不能下载；直接调用旧下载 URL 也被拒绝 ----------
def test_new_draft_cannot_be_downloaded(api, auth, h):
    d = new_draft(api, h)
    did = d["draft_id"]
    assert d["review_status"] == "draft" and "download_url" not in d
    for headers in (h, auth):                                    # 有无操作人都不放行
        r = api.get(f"/api/reports/{did}/download", headers=headers)
        assert r.status_code == 409 and r.json()["detail"]["error_code"] == "report_not_approved"
    # 初稿 DOCX 虽然落盘，但没有任何接口能取到它
    assert api.get(f"/api/reports/{did}/download", headers=h).headers["content-type"].startswith("application/json")


@pytest.mark.parametrize("path", [["start_review"], ["start_review", "request_changes"]])
def test_unfinished_states_cannot_be_downloaded(api, h, path):
    did = new_draft(api, h)["draft_id"]
    for a in path:
        assert act(api, h, did, a, comment="需要补充说明").status_code == 200
    r = api.get(f"/api/reports/{did}/download", headers=h)
    assert r.status_code == 409 and r.json()["detail"]["error_code"] == "report_not_approved"


def test_forged_reviewed_status_without_approval_record_is_not_exported(api, h):
    """即使有人绕过接口把 review_status 直接改成 reviewed，没有通过记录也不能导出。"""
    did = new_draft(api, h)["draft_id"]
    with api.app.state.session_factory() as s:
        s.get(ReportDraft, did).review_status = "reviewed"
        s.commit()
    r = api.get(f"/api/reports/{did}/download", headers=h)
    assert r.status_code == 409 and r.json()["detail"]["error_code"] == "report_not_approved"


def test_review_actions_require_operator(api, auth):
    did = new_draft(api, auth)["draft_id"]
    for headers in (auth, {**auth, "X-Operator": "anonymous"}, {**auth, "X-Operator": "  "}):
        r = act(api, headers, did, "start_review")
        assert r.status_code == 400 and r.json()["detail"]["error_code"] == "operator_required"
    assert detail(api, auth, did)["review_status"] == "draft"


def test_review_endpoints_require_token(api):
    assert api.get("/api/reports").status_code == 401
    assert api.get("/api/reports/x").status_code == 401
    assert api.post("/api/reports/x/transition", json={"action": "approve"}).status_code == 401


# ---------- 2. 状态转换 ----------
def test_state_machine_full_cycle_and_invalid_transitions(api, h):
    d = new_draft(api, h)
    did = d["draft_id"]
    for bad in ("approve", "resubmit", "request_changes"):
        r = act(api, h, did, bad, comment="x")
        assert r.status_code == 409 and r.json()["detail"]["error_code"] == "invalid_transition", bad
    assert act(api, h, did, "start_review").json()["review_status"] == "in_review"
    assert act(api, h, did, "start_review").status_code == 409           # 不能重复开始
    r = act(api, h, did, "request_changes")                              # 退回必须写意见
    assert r.status_code == 422 and r.json()["detail"]["error_code"] == "comment_required"
    assert act(api, h, did, "request_changes", comment="第 2 章措辞需调整").json()["review_status"] == "changes_requested"
    assert act(api, h, did, "approve", **approve_kw(d)).status_code == 409   # 退回后不能直接通过，须重新提交
    assert act(api, h, did, "resubmit").json()["review_status"] == "in_review"
    ok = act(api, h, did, "approve", **approve_kw(d))
    assert ok.status_code == 200 and ok.json()["review_status"] == "reviewed"
    assert act(api, h, did, "start_review").status_code == 409           # 终态

    ev = api.get(f"/api/reports/{did}/events", headers=h).json()
    assert [e["action"] for e in ev] == ["generated", "start_review", "request_changes", "resubmit", "approve"]
    assert [(e["from_status"], e["to_status"]) for e in ev[1:]] == [
        ("draft", "in_review"), ("in_review", "changes_requested"), ("changes_requested", "in_review"),
        ("in_review", "reviewed")]
    assert all(e["operator"] == "reviewer-a" and e["ts"] for e in ev[1:])
    assert ev[2]["comment"] == "第 2 章措辞需调整"


def test_list_filters_by_status(api, h):
    a, b = new_draft(api, h)["draft_id"], new_draft(api, h, "PRJ-BIO-002", "sample_analysis")["draft_id"]
    to_in_review(api, h, b)
    lst = api.get("/api/reports", headers=h).json()
    assert {x["draft_id"] for x in lst} >= {a, b}
    assert [x["draft_id"] for x in api.get("/api/reports", headers=h, params={"status": "in_review"}).json()] == [b]
    assert api.get("/api/reports", headers=h, params={"status": "bogus"}).status_code == 422


# ---------- 3. 通过前检查 ----------
def test_approve_requires_opinion_confirmations_and_missing_field_notes(api, h):
    d = new_draft(api, h, "PRJ-SM-003", "sample_analysis")
    did = d["draft_id"]
    assert d["missing_fields"]
    to_in_review(api, h, did)
    kw = approve_kw(d)

    r = act(api, h, did, "approve", **{**kw, "comment": "  "})
    assert r.status_code == 422 and r.json()["detail"]["error_code"] == "comment_required"
    r = act(api, h, did, "approve", **{**kw, "confirmations": ["facts_checked"]})
    assert r.status_code == 422 and r.json()["detail"]["error_code"] == "confirmations_required"
    r = act(api, h, did, "approve", **{**kw, "missing_ack": {}})                       # 缺失项不能靠勾选带过
    assert r.status_code == 422 and r.json()["detail"]["error_code"] == "missing_fields_unacknowledged"
    partial = dict(list(kw["missing_ack"].items())[:1])
    r = act(api, h, did, "approve", **{**kw, "missing_ack": partial})
    assert r.status_code == 422 and len(d["missing_fields"]) > 1
    assert detail(api, h, did)["review_status"] == "in_review"                       # 失败不改变状态

    assert act(api, h, did, "approve", **kw).status_code == 200
    ev = api.get(f"/api/reports/{did}/events", headers=h).json()[-1]
    assert set(ev["detail"]["missing_ack"]) == set(d["missing_fields"])


def test_open_issue_blocks_approval_until_resolved(api, h):
    d = new_draft(api, h)
    did = d["draft_id"]
    to_in_review(api, h, did)
    blk = pick(d, kind="fact", tagged=True)
    c = api.post(f"/api/reports/{did}/comments", headers=h,
                 json={"kind": "issue", "text": "该结论需与原始记录核对", "block_id": blk["bid"]})
    assert c.status_code == 200 and c.json()["status"] == "open"
    dd = detail(api, h, did)
    assert dd["readiness"]["can_approve"] is False
    assert dd["readiness"]["blockers"][0]["code"] == "open_issues"
    r = act(api, h, did, "approve", **approve_kw(d))
    assert r.status_code == 409 and r.json()["detail"]["error_code"] == "approval_blocked"

    assert api.post(f"/api/reports/{did}/comments/{c.json()['id']}/resolve", headers=h, json={}).status_code == 422
    res = api.post(f"/api/reports/{did}/comments/{c.json()['id']}/resolve", headers=h, json={"resolution": "已核对一致"})
    assert res.status_code == 200 and res.json()["resolved_by"] == "reviewer-a"
    assert act(api, h, did, "approve", **approve_kw(d)).status_code == 200


def test_plain_comment_does_not_block(api, h):
    d = new_draft(api, h)
    did = d["draft_id"]
    to_in_review(api, h, did)
    assert api.post(f"/api/reports/{did}/comments", headers=h, json={"kind": "comment", "text": "措辞可再精简"}).status_code == 200
    assert act(api, h, did, "approve", **approve_kw(d)).status_code == 200


def test_edit_that_loses_source_is_flagged_and_blocks_approval(api, h):
    d = new_draft(api, h)
    did = d["draft_id"]
    to_in_review(api, h, did)
    blk = pick(d, kind="fact", tagged=True, with_digits=True)
    bad_text = blk["text"] + " 另有 99999 例未列入。"                      # 新数字不在任何来源事实里
    r = api.post(f"/api/reports/{did}/blocks/{blk['bid']}/edit", headers=h,
                 json={"base_version": 1, "text": bad_text, "note": "补充说明"})
    assert r.status_code == 200 and r.json()["flag"] and "99999" in r.json()["flag"]
    dd = detail(api, h, did)
    flagged = [b for b in blocks(dd) if b.get("flag")]
    assert [b["bid"] for b in flagged] == [blk["bid"]] and dd["current_version"] == 2
    assert dd["readiness"]["blockers"][0]["code"] == "unsourced_edit"
    r = act(api, h, did, "approve", **approve_kw(d))
    assert r.status_code == 409 and r.json()["detail"]["error_code"] == "approval_blocked"

    # 改回有来源的写法后，标记消除，可以通过；原文保留
    good = blk["text"] + "（审核补充：与原始记录核对一致）"
    r = api.post(f"/api/reports/{did}/blocks/{blk['bid']}/edit", headers=h,
                 json={"base_version": 2, "text": good, "note": "去掉无来源数字"})
    assert r.status_code == 200 and r.json()["flag"] is None
    nb = next(b for b in blocks(detail(api, h, did)) if b["bid"] == blk["bid"])
    assert nb["edited"] and nb["original_text"] == blk["text"] and nb["tags"] == blk["tags"]
    assert act(api, h, did, "approve", **approve_kw(d)).status_code == 200


def test_new_business_conclusion_without_source_is_flagged(api, h):
    d = new_draft(api, h)
    did = d["draft_id"]
    to_in_review(api, h, did)
    anchor = pick(d, kind="fact", tagged=True)
    r = api.post(f"/api/reports/{did}/blocks", headers=h, json={
        "base_version": 1, "after_block": anchor["bid"], "kind": "fact", "text": "本批所有样品均符合要求。",
        "fact_keys": [], "note": "补充结论"})
    assert r.status_code == 200 and r.json()["flag"]
    bad_bid = r.json()["bid"]
    assert act(api, h, did, "approve", **approve_kw(d)).status_code == 409
    # 另加一条合规的说明性文字不能抵消无来源结论
    r = api.post(f"/api/reports/{did}/blocks", headers=h, json={
        "base_version": 2, "after_block": anchor["bid"], "kind": "static", "text": "审核人补充：详见附录。",
        "fact_keys": [], "note": "补充说明"})
    assert r.status_code == 200 and r.json()["flag"] is None
    blk = detail(api, h, did)["readiness"]["blockers"][0]
    assert blk["code"] == "unsourced_edit" and blk["block_ids"] == [bad_bid]
    # 新增位置在锚点之后；不可在附录 / 签名章节新增
    sig = next(b for b in blocks(detail(api, h, did)) if "签名" in b["section"] and b["type"] == "sentence")
    r = api.post(f"/api/reports/{did}/blocks", headers=h, json={
        "base_version": 3, "after_block": sig["bid"], "kind": "static", "text": "已签署。", "fact_keys": [], "note": "x"})
    assert r.status_code == 422 and r.json()["detail"]["error_code"] == "section_locked"


def test_static_text_cannot_gain_numbers_and_placeholders_cannot_be_removed(api, h):
    d = new_draft(api, h, "PRJ-SM-003", "sample_analysis")
    did = d["draft_id"]
    to_in_review(api, h, did)
    st = pick(d, kind="static")
    r = api.post(f"/api/reports/{did}/blocks/{st['bid']}/edit", headers=h,
                 json={"base_version": 1, "text": st["text"] + "合格率 95%。", "note": "x"})
    assert r.status_code == 200 and r.json()["flag"]                      # 说明性文字不得含数字 -> 标记

    fs = next(b for b in blocks(d) if b["type"] == "sentence" and b["kind"] == "fact" and "[待补充:" in b["text"])
    r = api.post(f"/api/reports/{did}/blocks/{fs['bid']}/edit", headers=h,
                 json={"base_version": 2, "text": "已补齐。", "fact_keys": fs["fact_keys"], "note": "补齐"})
    assert r.status_code == 422 and r.json()["detail"]["error_code"] == "placeholder_removed"


def test_missing_value_placeholders_cannot_be_edited_away(api, h):
    from sqlalchemy import update

    from app.db.models import Project, ReportVersion
    with api.app.state.session_factory() as s:
        s.execute(update(Project).where(Project.id == "PRJ-SM-001").values(sample_prep=None))
        s.commit()
    d = new_draft(api, h)
    did = d["draft_id"]
    ph = next(b for b in blocks(d) if b["type"] == "sentence" and "[待补充: 样品前处理方法]" in b["text"])
    r = api.post(f"/api/reports/{did}/blocks/{ph['bid']}/edit", headers=h,
                 json={"base_version": 1, "text": "样品前处理方法：蛋白沉淀法", "fact_keys": ph["fact_keys"],
                       "note": "补齐"})                                        # 不能靠改文字伪装已补齐
    assert r.status_code == 422 and r.json()["detail"]["error_code"] == "placeholder_removed"
    assert detail(api, h, did)["current_version"] == 1

    # kind=placeholder 的句子（builder.todo 产生）整句不可编辑：把该句的类型改为 placeholder 后验证
    with api.app.state.session_factory() as s:
        v = s.get(ReportVersion, (did, 1))
        import copy
        content = copy.deepcopy(v.content_json)
        for sec in content["sections"]:
            for blk in sec["blocks"]:
                if blk["bid"] == ph["bid"]:
                    blk["kind"], blk["fact_keys"] = "placeholder", []
        v.content_json = content
        s.commit()
    r = api.post(f"/api/reports/{did}/blocks/{ph['bid']}/edit", headers=h,
                 json={"base_version": 1, "text": ph["text"] + "。", "note": "x"})
    assert r.status_code == 422 and r.json()["detail"]["error_code"] == "not_editable"


def test_tables_and_unknown_blocks_are_not_editable(api, h):
    d = new_draft(api, h)
    did = d["draft_id"]
    tb = next(b for b in blocks(d) if b["type"] == "table")
    assert tb["editable"] is False
    r = api.post(f"/api/reports/{did}/blocks/{tb['bid']}/edit", headers=h, json={"base_version": 1, "text": "x", "note": "x"})
    assert r.status_code == 422 and r.json()["detail"]["error_code"] == "not_editable"
    r = api.post(f"/api/reports/{did}/blocks/B999/edit", headers=h, json={"base_version": 1, "text": "x", "note": "x"})
    assert r.status_code == 404
    st = pick(d, kind="static")
    for body, code in [({"text": "", "note": "x"}, "invalid_text"), ({"text": "新文字", "note": " "}, "note_required"),
                       ({"text": st["text"], "note": "x"}, "no_change")]:
        r = api.post(f"/api/reports/{did}/blocks/{st['bid']}/edit", headers=h, json={"base_version": 1, **body})
        assert r.status_code == 422 and r.json()["detail"]["error_code"] == code, (body, r.text)
    assert detail(api, h, did)["current_version"] == 1                     # 失败的修订不产生新版本


def test_stale_version_and_locked_report(api, h):
    d = new_draft(api, h)
    did = d["draft_id"]
    st = pick(d, kind="static")
    ok = api.post(f"/api/reports/{did}/blocks/{st['bid']}/edit", headers=h,
                  json={"base_version": 1, "text": st["text"] + "（修订）", "note": "措辞"})
    assert ok.status_code == 200
    stale = api.post(f"/api/reports/{did}/blocks/{st['bid']}/edit", headers=h,
                     json={"base_version": 1, "text": "另一处修订", "note": "措辞"})
    assert stale.status_code == 409 and stale.json()["detail"]["error_code"] == "version_conflict"
    to_in_review(api, h, did)
    assert act(api, h, did, "approve", **approve_kw(d)).status_code == 200
    r = api.post(f"/api/reports/{did}/blocks/{st['bid']}/edit", headers=h,
                 json={"base_version": 2, "text": "通过后再改", "note": "x"})
    assert r.status_code == 409 and r.json()["detail"]["error_code"] == "report_locked"
    r = api.post(f"/api/reports/{did}/comments", headers=h, json={"kind": "comment", "text": "晚了"})
    assert r.status_code == 409


# ---------- 4. 导出的是最终审核版本，且可由 python-docx 打开 ----------
def test_export_reflects_final_reviewed_content(api, h):
    d = new_draft(api, h, "PRJ-SM-003", "sample_analysis")
    did = d["draft_id"]
    to_in_review(api, h, did)
    st = pick(d, kind="static")
    new_text = st["text"] + "（审核人补充：本节仅为描述性说明）"
    assert api.post(f"/api/reports/{did}/blocks/{st['bid']}/edit", headers=h,
                    json={"base_version": 1, "text": new_text, "note": "补充说明"}).status_code == 200
    anchor = pick(d, kind="fact", tagged=True)
    added = "审核人新增：详见数据溯源附录。"
    assert api.post(f"/api/reports/{did}/blocks", headers=h, json={
        "base_version": 2, "after_block": anchor["bid"], "kind": "static", "text": added, "fact_keys": [],
        "note": "指引"}).status_code == 200
    assert act(api, h, did, "approve", **approve_kw(d)).status_code == 200

    r = api.get(f"/api/reports/{did}/download", headers=h)
    assert r.status_code == 200 and "wordprocessingml" in r.headers["content-type"]
    assert f"{did}-v3-reviewed.docx" in r.headers["content-disposition"]
    doc = Document(io.BytesIO(r.content))                                  # 可被 python-docx 打开
    paras = [p.text for p in doc.paragraphs]
    joined = "\n".join(paras)
    assert paras[0] == "AI 辅助草稿，需人工复核"                                # 文首标记仍在
    assert any(new_text in p and p.endswith("[人工修订]") for p in paras)       # 最终修订内容，而非初稿
    assert not any(p == st["text"] for p in paras)                          # 初稿原句不再出现
    assert any(added in p and p.endswith("[人工新增]") for p in paras)
    assert "人工审核记录" in joined and "reviewer-a" in joined and "审核版本：v3" in joined
    assert "不是受监管的电子签名" in joined and "不等于正式 QA 批准" in joined
    assert "[待补充: 仪器型号与色谱柱]" in joined                              # 缺失项仍是占位，没有被伪装成已补齐
    for f in d["missing_fields"]:
        assert f"待补充字段“{f}”" in joined
    assert "模拟数据" in doc.sections[0].header.paragraphs[0].text
    sig = next(t for t in doc.tables if [c.text for c in t.rows[0].cells] == ["角色", "姓名", "签名", "日期"])
    assert all(row.cells[1].text == "[待签署]" for row in sig.rows[1:])         # 审核通过不等于签署

    ev = api.get(f"/api/reports/{did}/events", headers=h).json()
    assert ev[-1]["action"] == "export" and ev[-1]["detail"]["version"] == 3 and ev[-1]["operator"] == "reviewer-a"
    assert api.get(f"/api/reports/{did}/download", headers={**h, "X-Operator": ""}).status_code == 400   # 导出也要操作人记录


def test_exported_docx_numbers_still_trace_to_facts(api, h):
    """导出内容经审核版本重新渲染后，所有带数字的叙述句仍带来源标签，且标签都在附录里。"""
    import re
    d = new_draft(api, h)
    did = d["draft_id"]
    to_in_review(api, h, did)
    assert act(api, h, did, "approve", **approve_kw(d)).status_code == 200
    doc = Document(io.BytesIO(api.get(f"/api/reports/{did}/download", headers=h).content))
    appendix = {r.cells[0].text for t in doc.tables if t.rows[0].cells[0].text == "编号" for r in t.rows[1:]}
    tagged = [re.search(r"\[(F\d+(?:,F\d+)*)\]", p.text) for p in doc.paragraphs]
    tags = {t for m in tagged if m for t in m.group(1).split(",")}
    assert tags and tags <= appendix


# ---------- 5. 刷新后继续审核；历史版本与审核记录可查询 ----------
def test_progress_survives_restart_and_history_is_queryable(api, h):
    d = new_draft(api, h)
    did = d["draft_id"]
    to_in_review(api, h, did)
    st = pick(d, kind="static")
    api.post(f"/api/reports/{did}/blocks/{st['bid']}/edit", headers=h,
             json={"base_version": 1, "text": st["text"] + "（改）", "note": "措辞"})
    cid = api.post(f"/api/reports/{did}/comments", headers=h, json={"kind": "issue", "text": "待核对"}).json()["id"]

    fresh = TestClient(create_app(api.app.state.settings))                  # 模拟服务重启 / 页面刷新后重新连接
    again = detail(fresh, h, did)
    assert again["review_status"] == "in_review" and again["current_version"] == 2
    assert [c["id"] for c in again["comments"]] == [cid] and again["comments"][0]["status"] == "open"
    assert any(x["draft_id"] == did and x["review_status"] == "in_review" and x["open_issues"] == 1
               for x in fresh.get("/api/reports", headers=h).json())

    vs = fresh.get(f"/api/reports/{did}/versions", headers=h).json()
    assert [(v["version"], v["change_type"], v["is_current"]) for v in vs] == [(1, "generated", False), (2, "edit", True)]
    v1 = detail(fresh, h, did, version=1)                                   # 历史版本仍是生成时原文
    assert v1["is_current"] is False and next(b for b in blocks(v1) if b["bid"] == st["bid"])["text"] == st["text"]
    assert next(b for b in blocks(again) if b["bid"] == st["bid"])["text"].endswith("（改）")
    assert fresh.get(f"/api/reports/{did}", headers=h, params={"version": 9}).status_code == 404
    actions = [e["action"] for e in fresh.get(f"/api/reports/{did}/events", headers=h).json()]
    assert actions == ["generated", "start_review", "edit", "issue_added"]


def test_review_actions_are_audited(api, h):
    d = new_draft(api, h)
    did = d["draft_id"]
    r = act(api, h, did, "start_review")
    with api.app.state.session_factory() as s:
        rows = s.scalars(select(AuditLog).where(AuditLog.request_id == r.headers["x-request-id"])).all()
    assert len(rows) == 1 and rows[0].draft_id == did and rows[0].operator == "reviewer-a" and rows[0].status_code == 200
    assert "start_review" in rows[0].input_digest or rows[0].feature.endswith("/transition")


# ---------- 旧数据兼容 ----------
def test_legacy_drafts_are_not_reviewable_and_never_look_reviewed(tmp_path):
    settings = Settings(database_url=f"sqlite:///{(tmp_path / 'old.db').as_posix()}", api_token="test-token",
                        audit_jsonl_dir=str(tmp_path / "audit"), reports_dir=str(tmp_path / "reports"))
    app = create_app(settings)
    from seed.loader import seed_database
    with app.state.session_factory() as s:
        seed_database(s)
        s.add(ReportDraft(id="DRAFT-VAL-OLD-000001", project_id="PRJ-SM-001", run_ids=[], kind="validation",
                          path=str(tmp_path / "old.docx"), review_status="reviewed", facts_json=[], missing_fields=[]))
        s.commit()
    (tmp_path / "old.docx").write_bytes(b"old")
    app = create_app(settings)                                              # 重新启动：标准化旧草稿
    c = TestClient(app)
    hh = {"Authorization": "Bearer test-token", "X-Operator": "reviewer-a"}
    row = c.get("/api/reports/DRAFT-VAL-OLD-000001", headers=hh).json()
    assert row["legacy"] is True and row["review_status"] == "draft" and "sections" not in row
    for a in ("start_review", "approve"):
        r = act(c, hh, "DRAFT-VAL-OLD-000001", a, comment="x")
        assert r.status_code == 409 and r.json()["detail"]["error_code"] == "legacy_draft"
    r = c.get("/api/reports/DRAFT-VAL-OLD-000001/download", headers=hh)
    assert r.status_code == 409 and r.json()["detail"]["error_code"] == "legacy_draft"


def test_old_database_without_review_columns_is_upgraded(tmp_path):
    """旧库的 report_draft 表没有 current_version 列：启动时补列，旧记录可读取。"""
    import sqlalchemy as sa
    url = f"sqlite:///{(tmp_path / 'older.db').as_posix()}"
    eng = sa.create_engine(url)
    with eng.begin() as conn:
        conn.execute(text("CREATE TABLE report_draft (id VARCHAR(40) PRIMARY KEY, project_id VARCHAR(32), run_ids JSON, "
                          "kind VARCHAR(20), path VARCHAR(300), created_at DATETIME, review_status VARCHAR(10), "
                          "facts_json JSON, missing_fields JSON)"))
        conn.execute(text("INSERT INTO report_draft VALUES ('DRAFT-VAL-X-1','PRJ-SM-001','[]','validation','x.docx',"
                          "'2026-01-01 00:00:00','draft','[]','[]')"))
    eng.dispose()
    app = create_app(Settings(database_url=url, api_token="test-token", audit_jsonl_dir=str(tmp_path / "a"),
                              reports_dir=str(tmp_path / "r")))
    r = TestClient(app).get("/api/reports/DRAFT-VAL-X-1", headers={"Authorization": "Bearer test-token"})
    assert r.status_code == 200 and r.json()["legacy"] is True
