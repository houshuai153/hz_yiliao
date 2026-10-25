"""三个评测套件。每个套件返回 (指标值字典, 逐题明细)。指标 = {name: (value, numerator, denominator)}。

评测只依据系统的**输出**（答案 / 引用 / 翻译结果 / 生成的 DOCX 文件）和对数据库的**独立查询**，
不调用被评测模块内部的校验函数来判断自己。
"""
import json
import re
from pathlib import Path

from docx import Document
from sqlalchemy import func, select, update

from app.db.models import CalibrationPoint, ISRRecord, KBSection, QCSample, StudySample
from app.evaluation.env import full_env, seeded_db
from app.kb.indexer import load_index
from app.llm.base import DeterministicProvider
from app.regulatory.config import load_regulatory_config
from app.regulatory.service import RegulatoryAssistant
from app.reports.service import generate_draft
from app.rules.engine import load_rule_sets
from app.translation.service import TranslationService

Metric = tuple  # (value, numerator, denominator)


def rate(num: int, den: int) -> Metric:
    return (num / den if den else None, num, den)


def squash(s: str) -> str:
    return re.sub(r"\s+", "", s)


def load_cases(path: str | Path) -> list[dict]:
    return json.loads(Path(path).read_text(encoding="utf-8"))["cases"]


# ============================ 法规问答 ============================
DISCLAIMER_PHRASE = "非监管机构原文"


def run_regulatory(cases: list[dict]) -> tuple[dict[str, Metric], list[dict]]:
    with full_env() as env:
        assistant = RegulatoryAssistant(load_index(env.settings.index_dir), load_regulatory_config())
        with env.session_factory() as s:      # 独立于助手内部缓存，直接取库中章节原文与元数据
            rows = {r.section_id: r for r in s.scalars(select(KBSection))}
        details, c = [], dict(hit=0, false_ref=0, ans_n=0, rec_ok=0, ref_n=0, cit_ok=0, cit_n=0,
                              lab_ok=0, lab_n=0, prec_ok=0, prec_n=0)
        for case in cases:
            r = assistant.ask(case["q"])
            cited = [x["section_id"] for x in r.citations]
            d = {"id": case["id"], "question": case["q"], "kind": case["kind"], "refused": r.refused,
                 "refusal_reason": r.refusal_reason, "confidence": r.confidence, "cited_sections": cited}
            if case["kind"] == "answerable":
                c["ans_n"] += 1
                expect = set(case["expect_sections"])
                hit = (not r.refused) and bool(expect & set(cited))
                c["hit"] += hit
                c["false_ref"] += r.refused
                c["prec_n"] += len(cited)
                c["prec_ok"] += sum(x in expect for x in cited)
                d.update(expect_sections=sorted(expect), passed=hit)
            else:
                c["ref_n"] += 1
                ok = r.refused and (case.get("expect_reason") in (None, r.refusal_reason))
                c["rec_ok"] += ok
                d.update(expect_reason=case.get("expect_reason"), passed=ok)
            if not r.refused:
                for x in r.citations:
                    c["cit_n"] += 1
                    row = rows.get(x["section_id"])
                    c["cit_ok"] += bool(row and squash(x["quote"]) in squash(row.text)
                                        and (x["doc_id"], x["version"], x["source_type"]) == (row.doc_id, row.version, row.source_type)
                                        and x["is_demo"] is True and row.is_demo)
                c["lab_n"] += 1
                c["lab_ok"] += (DISCLAIMER_PHRASE in r.answer) and bool(r.citations) and all(
                    x["is_demo"] and ("非监管原文" in x["label"] or "无合规效力" in x["label"]) for x in r.citations)
                d["answer_preview"] = r.answer[:120]
            details.append(d)
    return {
        "answerable_hit_rate": rate(c["hit"], c["ans_n"]),
        "false_refusal_rate": rate(c["false_ref"], c["ans_n"]),
        "refusal_recall": rate(c["rec_ok"], c["ref_n"]),
        "citation_validity": rate(c["cit_ok"], c["cit_n"]),
        "demo_labeling_rate": rate(c["lab_ok"], c["lab_n"]),
        "citation_precision": rate(c["prec_ok"], c["prec_n"]),
    }, details


# ============================ 翻译与术语核对 ============================
def run_translation(cases: list[dict]) -> tuple[dict[str, Metric], list[dict]]:
    provider = DeterministicProvider()
    svc = TranslationService(provider, provider.glossary)
    c = dict(status_ok=0, n=0, exact_ok=0, exact_n=0, term_ok=0, term_n=0, warn_ok=0, warn_n=0,
             unexpected=0, nofab_ok=0, nofab_n=0)
    details = []
    for case in cases:
        res = svc.translate(case["text"], case["direction"], case.get("reference"))
        p, exp = res["paragraphs"][0], case["expect"]
        hit_terms = {h["term_id"] for h in p["term_hits"]}
        got_codes = {w["code"] for w in p["warnings"]}
        c["n"] += 1
        c["status_ok"] += p["status"] == exp["status"]
        if "target" in exp:
            c["exact_n"] += 1
            c["exact_ok"] += p["target"] == exp["target"]
        c["term_n"] += len(exp["term_ids"])
        c["term_ok"] += sum(t in hit_terms for t in exp["term_ids"])
        c["warn_n"] += len(exp["warning_codes"])
        c["warn_ok"] += sum(w in got_codes for w in exp["warning_codes"])
        unexpected = [w["code"] for w in p["warnings"] if w["severity"] == "error" and w["code"] not in exp["warning_codes"]]
        c["unexpected"] += len(unexpected)
        if exp["status"] == "unsupported":
            c["nofab_n"] += 1
            c["nofab_ok"] += p["target"] is None and not any(s["status"] == "translated" for s in p["sentences"])
        details.append({"id": case["id"], "status": p["status"], "target": p["target"],
                        "term_hits": sorted(hit_terms), "warning_codes": sorted(got_codes),
                        "unexpected_errors": unexpected,
                        "passed": p["status"] == exp["status"] and all(w in got_codes for w in exp["warning_codes"])
                        and all(t in hit_terms for t in exp["term_ids"]) and not unexpected
                        and ("target" not in exp or p["target"] == exp["target"])})
    return {
        "status_accuracy": rate(c["status_ok"], c["n"]),
        "exact_match_rate": rate(c["exact_ok"], c["exact_n"]),
        "term_hit_recall": rate(c["term_ok"], c["term_n"]),
        "warning_recall": rate(c["warn_ok"], c["warn_n"]),
        "unexpected_error_warnings": (c["unexpected"], c["unexpected"], None),
        "no_fabrication_rate": rate(c["nofab_ok"], c["nofab_n"]),
    }, details


# ============================ 报告 ============================
REQUIRED_HEADINGS = ["方法概述", "校准曲线", "质控样品", "分析批接受判定", "已测样品再分析", "偏差", "结论",
                     "复核清单", "数据溯源附录", "签名"]
KIND_HEADING = {"validation": "准确度与精密度", "sample_analysis": "研究样品分析概况"}
BANNER = "AI 辅助草稿，需人工复核"
_TAG_RE = re.compile(r"\s\[(F\d+(?:,F\d+)*)\]$")
_PH_RE = re.compile(r"\[待补充: [^\]]*\]")
_ID_RE = re.compile(r"[A-Z][A-Z0-9]*(?:-[A-Z0-9]+)+")
_NUM_RE = re.compile(r"\d+(?:\.\d+)?")
FORBIDDEN_TOKENS = ("None", "nan", "NaN", "null")


def _numbers(text: str) -> list[str]:
    return _NUM_RE.findall(_ID_RE.sub("", _PH_RE.sub("", _TAG_RE.sub("", text))))


def _all_text(doc) -> list[str]:
    out = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            out.extend(c.text for c in row.cells)
    return out


def _appendix(doc) -> dict[str, tuple[str, str]]:
    for t in doc.tables:
        head = [c.text for c in t.rows[0].cells]
        if head[:2] == ["编号", "事实键"]:
            return {r.cells[0].text: (r.cells[1].text, r.cells[3].text) for r in t.rows[1:]}
    return {}


def _integrity_ok(doc, kind: str) -> bool:
    paras = [p for p in doc.paragraphs if p.text.strip()]
    heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    sig = [t for t in doc.tables if [c.text for c in t.rows[0].cells] == ["角色", "姓名", "签名", "日期"]]
    return bool(
        paras and paras[0].text == BANNER
        and BANNER in doc.sections[0].header.paragraphs[0].text
        and all(any(h in x for x in heads) for h in REQUIRED_HEADINGS + [KIND_HEADING[kind]])
        and sig and [r.cells[0].text for r in sig[0].rows[1:]] == ["编写人", "复核人", "专题负责人", "质量保证"])


def run_report(cases: list[dict]) -> tuple[dict[str, Metric], list[dict]]:
    rule_sets = load_rule_sets()
    c = dict(integ_ok=0, n=0, cov_ok=0, cov_n=0, num_ok=0, num_n=0, agree_ok=0, agree_n=0, miss_ok=0)
    details = []
    for case in cases:
        with seeded_db() as (factory, reports_dir):
            with factory() as s:
                mut = case.get("mutation", {}).get("null_qc")
                if mut:
                    s.execute(update(QCSample).where(QCSample.run_id == mut["run_id"], QCSample.level == mut["level"])
                              .values(measured=None, bias_pct=None))
                    s.commit()
                res = generate_draft(s, reports_dir, case["project_id"], case["run_ids"], case["kind"], rule_sets)
                doc = Document(res["path"])                    # 评测对象是落盘后重新打开的文件
                integ = _integrity_ok(doc, case["kind"])
                appendix = _appendix(doc)

                # 来源覆盖与数字一致（只看 DOCX 正文）
                cov_n = cov_ok = num_n = num_ok = 0
                for p in doc.paragraphs:
                    t = p.text.strip()
                    if (not t or p.style.name.startswith(("Heading", "Title")) or t.startswith(("项目：", "数据来源：", "AI 辅助草稿", "模拟数据"))):
                        continue
                    nums = _numbers(t)
                    if not nums:
                        continue
                    m = _TAG_RE.search(p.text)
                    cov_n += 1
                    cov_ok += bool(m and all(tag in appendix for tag in m.group(1).split(",")))
                    if m and all(tag in appendix for tag in m.group(1).split(",")):
                        allowed = {n for tag in m.group(1).split(",") for n in _NUM_RE.findall(appendix[tag][1])}
                        num_n += 1
                        num_ok += all(n in allowed for n in nums)

                # 附录事实与数据库的独立核对
                agree_n = agree_ok = 0
                for _tag, (key, display) in appendix.items():   # 附录映射：编号 -> (事实键, 取值)
                    m = re.fullmatch(r"(.+-R\d+)\.(smp|qc|cal|isr)\.n", key)
                    if not m:
                        continue
                    run_id, kind_ = m.groups()
                    table = {"smp": StudySample, "qc": QCSample, "cal": CalibrationPoint, "isr": ISRRecord}[kind_]
                    truth = s.scalar(select(func.count()).select_from(table).where(table.run_id == run_id))
                    agree_n += 1
                    agree_ok += display == str(truth)

                # 缺失值处理
                exp, texts = case["expect"], _all_text(doc)
                checks = {
                    "missing_fields": sorted(res["missing_fields"]) == sorted(exp["missing_fields"]),
                    "placeholder_counts": all(sum(t.count(ph) for t in texts) == n
                                              for ph, n in exp.get("placeholder_counts", {}).items()),
                    "absent_text": not any(any(a in t for t in texts) for a in exp.get("absent_text", [])),
                    "present_text": all(any(a in t for t in texts) for a in exp.get("present_text", [])),
                    "no_forbidden_tokens": not any(tok in "\n".join(texts) for tok in FORBIDDEN_TOKENS),
                }
                pm = exp.get("paragraph_must_contain")
                if pm:
                    paras = [p.text for p in doc.paragraphs if p.text.startswith(pm["startswith"])]
                    checks["paragraph_must_contain"] = len(paras) == 1 and all(x in paras[0] for x in pm["contains"])
        c["n"] += 1
        c["integ_ok"] += integ
        c["cov_n"] += cov_n; c["cov_ok"] += cov_ok; c["num_n"] += num_n; c["num_ok"] += num_ok
        c["agree_n"] += agree_n; c["agree_ok"] += agree_ok
        miss_ok = all(checks.values())
        c["miss_ok"] += miss_ok
        details.append({"id": case["id"], "draft_id": res["draft_id"], "docx_integrity": integ,
                        "source_coverage": [cov_ok, cov_n], "numeric_consistency": [num_ok, num_n],
                        "source_value_agreement": [agree_ok, agree_n], "missing_checks": checks,
                        "missing_fields": res["missing_fields"], "passed": integ and miss_ok and cov_ok == cov_n
                        and num_ok == num_n and agree_ok == agree_n})
    return {
        "docx_integrity_rate": rate(c["integ_ok"], c["n"]),
        "source_coverage": rate(c["cov_ok"], c["cov_n"]),
        "numeric_consistency": rate(c["num_ok"], c["num_n"]),
        "source_value_agreement": rate(c["agree_ok"], c["agree_n"]),
        "missing_value_handling_rate": rate(c["miss_ok"], c["n"]),
    }, details


SUITES = {"regulatory": run_regulatory, "translation": run_translation, "report": run_report}
