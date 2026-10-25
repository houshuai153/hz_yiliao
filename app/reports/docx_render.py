"""DOCX 渲染（python-docx）。只负责排版，不做任何计算。"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.reports.model import ReportDoc, Sentence, TableBlock

BANNER = "AI 辅助草稿，需人工复核"
SIM_MARK = "模拟数据 / SIMULATED DATA"
REVIEW_NOTE = ("以上为系统内“操作人记录”，不是受监管的电子签名；人工审核通过不等于正式 QA 批准，"
               "正式报告须按质量体系完成签署。")
FONT_CN = "Microsoft YaHei"


def _font(run, size: float | None = None, bold: bool | None = None, color: RGBColor | None = None) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _table(doc: Document, tb: TableBlock) -> None:
    cap = doc.add_paragraph()
    _font(cap.add_run(tb.caption), 10, True)
    t = doc.add_table(rows=1, cols=len(tb.header))
    t.style = "Table Grid"
    small = 8 if len(tb.header) >= 6 else 9
    for i, h in enumerate(tb.header):
        c = t.rows[0].cells[i]
        c.text = ""
        _font(c.paragraphs[0].add_run(h), small, True)
    for row in tb.rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            _font(cells[i].paragraphs[0].add_run(str(v)), small)
    src = doc.add_paragraph()
    _font(src.add_run(f"数据来源：{tb.source}"), 8, color=RGBColor(0x66, 0x66, 0x66))


def _sentence(doc: Document, rep: ReportDoc, s: Sentence) -> None:
    p = doc.add_paragraph()
    _font(p.add_run(s.text), 10.5)
    if s.kind == "fact" and s.fact_keys:
        tags = ",".join(rep.facts.get(k).tag for k in s.fact_keys)
        _font(p.add_run(f" [{tags}]"), 8, color=RGBColor(0x66, 0x66, 0x66))
    if s.edited:
        _font(p.add_run(" [人工新增]" if s.original_text is None else " [人工修订]"), 8, color=RGBColor(0x1F, 0x4E, 0x9A))


def _review_block(doc: Document, review: dict) -> None:
    doc.add_heading("人工审核记录", level=1)
    p = doc.add_paragraph()
    _font(p.add_run(f"审核版本：v{review['version']}　操作人记录：{review['operator']}　时间：{review['ts']}"), 10)
    p = doc.add_paragraph()
    _font(p.add_run(f"审核意见：{review['comment']}"), 10)
    for field_name, note in review["missing_ack"].items():
        p = doc.add_paragraph()
        _font(p.add_run(f"待补充字段“{field_name}”：审核时确认保留为待补充。处理说明：{note}"), 10)
    p = doc.add_paragraph()
    _font(p.add_run(REVIEW_NOTE), 9, True, RGBColor(0xC0, 0x00, 0x00))


def render_docx(rep: ReportDoc, path: str | Path, draft_id: str, generated_at: str, review: dict | None = None) -> Path:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)

    hdr = doc.sections[0].header.paragraphs[0]
    hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(hdr.add_run(f"{BANNER} ｜ {SIM_MARK}"), 9, True, RGBColor(0xC0, 0x00, 0x00))
    ftr = doc.sections[0].footer.paragraphs[0]
    ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    foot = f"草稿编号 {draft_id}" + (f" ｜ 已人工审核 v{review['version']}（非正式 QA 批准）" if review else "")
    _font(ftr.add_run(foot), 8, color=RGBColor(0x66, 0x66, 0x66))

    # 文首标记必须是正文第一段
    b = doc.add_paragraph()
    b.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(b.add_run(BANNER), 16, True, RGBColor(0xC0, 0x00, 0x00))
    m = doc.add_paragraph()
    m.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(m.add_run(f"{SIM_MARK}：本文件全部数据由原型系统生成，不含真实受试者信息"), 9, color=RGBColor(0xC0, 0x00, 0x00))

    if review:
        r = doc.add_paragraph()
        r.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(r.add_run(f"人工审核通过（版本 v{review['version']}）——仍为 AI 辅助草稿，需按正式流程签署"), 10, True,
              RGBColor(0x1F, 0x4E, 0x9A))
    doc.add_heading(rep.title, level=0)
    meta = doc.add_paragraph()
    _font(meta.add_run(f"项目：{rep.project_id}　纳入分析批：{'、'.join(rep.run_ids)}　生成时间：{generated_at}　草稿编号：{draft_id}"), 9)

    for sec in rep.sections:
        doc.add_heading(sec.title, level=1)
        for blk in sec.blocks:
            if isinstance(blk, TableBlock):
                _table(doc, blk)
            else:
                _sentence(doc, rep, blk)
    if review:
        _review_block(doc, review)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path
